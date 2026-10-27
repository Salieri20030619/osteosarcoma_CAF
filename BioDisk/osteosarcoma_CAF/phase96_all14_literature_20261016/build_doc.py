from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
O=Path(__file__).parent
c=json.loads((O/'content.json').read_text(encoding='utf-8'))
project=json.loads((O/'project_registry.json').read_text(encoding='utf-8'))
by={r['gene']:r for r in project};assert len(c['genes'])==14 and set(by)=={g['gene'] for g in c['genes']}
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Noto Sans CJK SC';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');s.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
 s=d.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
d.styles['Title'].paragraph_format.space_before=Pt(0)
d.styles['Subtitle'].font.italic=False
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.header.paragraphs[0].text='骨肉瘤 CAF 候选基因文献调查';sec.header.paragraphs[0].runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('研究讨论稿  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(8)
def p(text,boldlabel=None):
 x=d.add_paragraph()
 if boldlabel:x.add_run(boldlabel+'  ').bold=True
 x.add_run(text);return x
def heading(text,level=1):return d.add_heading(text,level)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   pr.append(borders);mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'60');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');pr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.13
    if ci==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs:r.font.size=Pt(9);r.bold=(ri==0)
 p('')
 return t
def hyperlink(para,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'244A73');rp.append(color);r.append(rp);text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);para._p.append(h)

d.add_paragraph('骨肉瘤CAF十四个候选基因\n文献调查与生物学解释','Title');d.add_paragraph(c['date']+'  文献检索与项目结果解释','Subtitle')
p('这14个基因可以组织成骨架与黏附、核膜力学、代谢与氧化还原、转录调控、膜相关信号及非编码RNA探索等研究方向。现有文献支持若干具体机制假设，其中EBF1与CAF问题的联系较直接；但它们不是14个已确认的骨肉瘤CAF驱动基因，也不是现成的活细胞流式标志组合。')
p('本文件逐一核对基因身份、原始文献模型、主要发现及对本项目的可用解释，保留关键阴性和相反结果。全部14个候选继续保留，文献多少不作为剔除标准。本轮新增的机制依据主要包括ZNF526的SHMT1相关研究、NDUFS7的基因缺失与抗氧化补偿研究，以及EBF1的其他癌种CAF功能研究。')
heading('项目证据定位')
p('候选名单与统计数字来自2026年10月12日冻结的Phase84研究登记，基于历史拓展层及按来源分别进行的患者层面骨病灶与肺病灶比较。SCPCP000017为骨12例、肺3例，SCPCP000023为骨4例、肺3例。13个候选在两来源均为肺端较高，EBF1则均为骨端较高；14个基因均有原始P＜0.05，均未达到FDR＜0.05，故保留“跨来源方向一致的名义显著探索性候选”定位。完整方向、效应、P值和FDR见附表。')
p('文献解释不替代统计检验。此前Phase86功能程序分析的500项有效检验没有原始P＜0.05的结果，因此当前数据不支持把逐基因信号直接写成已证实的整体功能程序增强。NCC新增资源和正在后台处理的19份NCI样本尚未用于更新本文件中的14基因统计结果。')
heading('检索与证据归类')
p('检索日期为2026年10月16日。围绕14个正式符号、项目稳定Ensembl ID及主要旧名，结合osteosarcoma、CAF、fibroblast、cancer和机制词定向检索；以PubMed／PMC、Europe PMC、原期刊和NCBI Gene为主要核验入口。检索为针对本项目的定向调查，不宣称穷尽全部文献。可读正文用于核对实验对象和结论，只有摘要可核验的研究不扩展其方法细节。')
p('文献按直接CAF研究、骨肉瘤肿瘤细胞研究、其他癌种功能实验、统计关联及基础注释／结构研究区分。讨论中的“本项目解释”和“建议验证”均为待检验的研究判断，不是文献已经证明的骨肉瘤CAF结论。参考文献提供可点击原文入口。')
heading('十四个候选的证据概览').paragraph_format.page_break_before=True
table(['基因','主要研究方向','最相关文献证据'],[(g['gene'],g['axis'],g['evidence']) for g in c['genes']],[2.4,5.0,9.4])
p('表中“功能实验”均限定于相应文献模型；它不自动代表本项目细胞已具备相同功能。“基础注释”也不代表无研究价值，而是提示先验证身份、表达和定位。')
heading('逐基因文献与生物学解释').paragraph_format.page_break_before=True
for i,g in enumerate(c['genes'],1):
 heading(str(i)+' '+g['gene']+' '+g['axis'],2)
 p(g['key']).paragraph_format.keep_with_next=True
 direction='骨端较高' if g['gene']=='EBF1' else '肺端较高'
 p('稳定标识 '+by[g['gene']]['feature_ids']+'；本项目两来源均'+direction+'；主要证据 '+g['evidence']).paragraph_format.keep_with_next=True
 for key,label in [('biology','生物学定位'),('literature','原始研究'),('interpretation','本项目解释'),('validation','建议验证')]:p(g[key],label)

heading('如何把十四个候选组织成可检验的生物学问题')
p('第一条线索是细胞对基质环境的适应。NEBL和RTKN偏向胞质骨架／黏附，TOR1AIP1偏向核膜适应；它们可以分别检验细胞在不同机械环境中的行为。当前证据不足以把三者连成一条已确立的信号通路，也不应由迁移相关文献推导出胶原沉积增加。')
p('第二条线索是代谢和氧化还原状态。NDUFS7与复合体I有关，ZNF526的SHMT1机制指向一碳代谢和谷胱甘肽，PNKD在肾癌模型中关联乳酸与PI3K–AKT–mTOR。可把这些作为并行假设，分别测量代谢通量和功能输出；不能把它们的转录本共同升高等同于同一种代谢模式。')
p('第三条线索是间质状态及微环境相互作用。EBF1既有CAF功能研究，也有正常骨髓及肺间质背景证据。本项目中EBF1为骨端较高，提示应关注骨相关间质身份或状态，而非统一解释成肺端激活。骨肺差异还可能包含组织常驻细胞组成的变化。区分“相同状态内表达改变”与“某类细胞比例改变”，需要在来源、患者和谱系分层后分析。[23–25]')
p('GRID2、BEST1、CALHM5与三个非编码RNA适合从表达定位和分子身份切入。对这些候选，清晰、可重复的细胞归属比先拼接一条复杂通路更有价值。验证中应独立识别CAF候选，避免用目标基因本身定义细胞群后，再把同一基因的富集视为独立证据。')
heading('蛋白定位与流式研究的关系',2)
table(['候选类别','涉及基因','可采用的解释与检测定位'],[
 ('胞内蛋白','NEBL、RTKN、NDUFS7、PNKD\nTOR1AIP1、TAF10\nZNF526、EBF1','适合蛋白、胞内或原位验证；不能直接作为活细胞表面分选标志。'),
 ('潜在膜相关候选','GRID2、BEST1、CALHM5','先核对目标细胞的表面表达、表位可及性及非透膜检测；膜注释不等于已有可用CAF流式抗体。'),
 ('非编码RNA','PEF1-AS1、RGS2-AS1\nAL096854.1','适合转录本／链特异性检测与RNA原位；没有对应蛋白可直接用于表面抗体分选。')],[2.8,5.4,8.6])
heading('验证顺序和推断边界',2)
p('先验证分子身份和细胞归属，再验证患者层面的表达方向，最后开展功能干预和回补。全部14个候选均保留；按文献成熟度安排实验顺序不改变冻结名单，也不重新计算原候选排名。')
p('新增NCI数据可以检验肺端表达是否存在、来自何种细胞状态，以及患者之间是否一致。只有肺端的数据不能单独构成骨肺组间差异复现；与外部骨端比较时，应明确非配对设计、来源与部位混杂，并与已有同来源比较核对方向。技术重复和重复取材不能增加独立患者N。')
p('骨肉瘤肿瘤细胞、其他癌种CAF、正常组织间质和本项目拓展候选是不同证据层。RNA变化不等于蛋白变化，表达相关不等于因果作用，通路名称不等于功能读数。主要相反或阴性结果已保留：NEBL及TOR1AIP1存在模型间方向差异，RGS2-AS1旧名对应研究出现阴性qPCR，AL096854.1未进入其文献最终模型。[20,21,26,27]')

heading('附表 原候选统计结果').paragraph_format.page_break_before=True
p('效应Δ为肺减骨，来自患者与部位层面汇总原始计数后计算的log2(CP10k＋0.1)表达均值差。它不是细胞均值log1p分数，也不能直接当作未经伪计数影响的表达倍数。患者标签精确置换检验按来源独立进行；两部位均有候选细胞的患者与非配对对比分开处理。')
headers=['基因','Δ S17','P S17','Δ S23','P S23']
rows=[]
for g in c['genes']:
 r=by[g['gene']];rows.append([g['gene'],f"{r['SCPCP000017__mean_lung_minus_bone']:.4f}",f"{r['SCPCP000017__P']:.6f}",f"{r['SCPCP000023__mean_lung_minus_bone']:.4f}",f"{r['SCPCP000023__P']:.6f}"])
table(headers,rows,[3.2,3.4,3.4,3.4,3.4])
q17={r['SCPCP000017__FDR'] for r in project};q23={r['SCPCP000023__FDR'] for r in project};assert len(q17)==len(q23)==1
p('S17＝SCPCP000017，骨N＝12、肺N＝3；S23＝SCPCP000023，骨N＝4、肺N＝3。效应展示至4位小数，P值展示至6位小数；原始完整精度保存在项目研究登记中。')
p('14个基因的S17 BH FDR均为 '+str(next(iter(q17)))+'；S23 BH FDR均为 '+str(next(iter(q23)))+'。BH检验家族分别为 all_cells__expanded__SCPCP000017__All_CAF__independent|genes 和 all_cells__expanded__SCPCP000023__All_CAF__independent|genes，未把不同来源、功能程序或文献结果混入同一个校正池。')
p('来源：phase84_all14_research_adoption_20261012 的 ALL14_NEXT_STAGE_RESEARCH_REGISTRY.tsv；效应定义回查phase81_allCAF_bone_lung_20261012代码。该附表是既有结果的回读，不是本轮重新检验。')

heading('参考文献与核验入口')
p('以下以原始研究和规范注释为主。文献年份采用出版记录；检索日期统一为2026年10月16日。补表记录仅用于说明关联线索，不替代基因功能研究。')
for r in c['references']:
 para=d.add_paragraph();para.paragraph_format.space_after=Pt(7);para.paragraph_format.line_spacing=1.12;para.paragraph_format.keep_together=True
 run=para.add_run(f"[{r['id']}] {r['title']}. {r['year']}. {r['identifier']}. ");run.font.size=Pt(9)
 hyperlink(para,'原文链接',r['url'])

for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
docpath=O/'骨肉瘤CAF十四个候选基因_文献调查与生物学解释_20261016.docx';d.save(docpath)
read=Document(docpath);text='\n'.join(p.text for p in read.paragraphs)+'\n'+'\n'.join(cell.text for t in read.tables for row in t.rows for cell in row.cells)
assert all(g['gene'] in text for g in c['genes']);assert len(c['references'])==27
assert '500项有效检验' in text and 'log2(CP10k＋0.1)' in text
for row in project:
 assert (row['SCPCP000017__mean_lung_minus_bone']<0)==(row['gene']=='EBF1')
 assert (row['SCPCP000023__mean_lung_minus_bone']<0)==(row['gene']=='EBF1')
(O/'BUILD_RECEIPT.json').write_text(json.dumps({'docx':str(docpath),'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'genes':14,'references':27,'tables':len(read.tables),'paragraphs':len(read.paragraphs),'render_verified':False},ensure_ascii=False,indent=2),encoding='utf-8')
print(docpath)
