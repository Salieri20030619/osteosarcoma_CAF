from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
O=Path(__file__).parent
c=json.loads((O/'content.json').read_text(encoding='utf-8'))
project=json.loads((O/'project_registry.json').read_text(encoding='utf-8'))
by={r['gene']:r for r in project};assert len(c['genes'])==14 and set(by)=={g['gene'] for g in c['genes']}
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Noto Sans CJK SC';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');s.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
 s=d.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
d.styles['Title'].paragraph_format.space_before=Pt(0)
d.styles['Subtitle'].font.italic=False
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.header.paragraphs[0].text='骨肉瘤 CAF 候选基因文献调查';sec.header.paragraphs[0].runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('研究讨论稿  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(8)
def p(text,boldlabel=None):
 x=d.add_paragraph()
 if boldlabel:x.add_run(boldlabel+'  ').bold=True
 x.add_run(text);return x
def heading(text,level=1):return d.add_heading(text,level)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   pr.append(borders);mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'60');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');pr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.13
    if ci==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs:r.font.size=Pt(9);r.bold=(ri==0)
 p('')
 return t
def hyperlink(para,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'244A73');rp.append(color);r.append(rp);text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);para._p.append(h)

d.add_paragraph('骨肉瘤CAF十四个候选基因\n文献调查与生物学解释','Title');d.add_paragraph(c['date']+'  文献检索与项目结果解释','Subtitle')
p('这14个基因可以组织成骨架与黏附、核膜力学、代谢与氧化还原、转录调控、膜相关信号及非编码RNA探索等研究方向。现有文献支持若干具体机制假设，其中EBF1与CAF问题的联系较直接；但它们不是14个已确认的骨肉瘤CAF驱动基因，也不是现成的活细胞流式标志组合。')
p('本文件逐一核对基因身份、原始文献模型、主要发现及对本项目的可用解释，保留关键阴性和相反结果。全部14个候选继续保留，文献多少不作为剔除标准。本轮新增的机制依据主要包括ZNF526的SHMT1相关研究、NDUFS7的基因缺失与抗氧化补偿研究，以及EBF1的其他癌种CAF功能研究。')
heading('项目证据定位')
p('候选名单与统计数字来自2026年10月12日冻结的Phase84研究登记，基于历史拓展层及按来源分别进行的患者层面骨病灶与肺病灶比较。SCPCP000017为骨12例、肺3例，SCPCP000023为骨4例、肺3例。13个候选在两来源均为肺端较高，EBF1则均为骨端较高；14个基因均有原始P＜0.05，均未达到FDR＜0.05，故保留“跨来源方向一致的名义显著探索性候选”定位。完整方向、效应、P值和FDR见附表。')
heading('检索与证据归类')
p('检索日期为2026年10月16日。围绕14个正式符号、项目稳定Ensembl ID及主要旧名，结合osteosarcoma、CAF、fibroblast、cancer和机制词定向检索；以PubMed／PMC、Europe PMC、原期刊和NCBI Gene为主要核验入口。检索为针对本项目的定向调查，不宣称穷尽全部文献。可读正文用于核对实验对象和结论，只有摘要可核验的研究不扩展其方法细节。')
p('文献按直接CAF研究、骨肉瘤肿瘤细胞研究、其他癌种功能实验、统计关联及基础注释／结构研究区分。讨论中的“本项目解释”和“建议验证”均为待检验的研究判断，不是文献已经证明的骨肉瘤CAF结论。参考文献提供可点击原文入口。')
heading('原始患者层面差异分析').paragraph_format.page_break_before=True
d.add_picture(str(O/'volcano.png'),width=Cm(16.7))
p('图1 两来源拓展层CAF的骨肺差异火山图。横轴为肺减骨的患者层面平均log2(CP10k＋0.1)表达差，纵轴为−log10（原始P值）；正值表示肺端较高。每个点对应一项有效基因检验，虚线为原始P＝0.05；该图不额外设置效应量筛选阈值。金色点标记14个保留候选，13个肺端较高、EBF1骨端较高。')
p('两来源分别进行患者标签精确置换检验，未合并样本或以细胞数替代患者数。颜色表示原始P值的名义显著性，FDR＜0.05数量另列于各面板；14个标记候选均未达到FDR＜0.05。完整检验中P值或效应量不可评估的条目不绘制。')

heading('十四个候选的证据概览').paragraph_format.page_break_before=True
table(['基因','主要研究方向','最相关文献证据'],[(g['gene'],g['axis'],g['evidence']) for g in c['genes']],[2.4,5.0,9.4])
p('表中“功能实验”均限定于相应文献模型；它不自动代表本项目细胞已具备相同功能。“基础注释”也不代表无研究价值，而是提示先验证身份、表达和定位。')
heading('逐基因文献与生物学解释').paragraph_format.page_break_before=True
for i,g in enumerate(c['genes'],1):
 heading(str(i)+' '+g['gene']+' '+g['axis'],2)
 p(g['key']).paragraph_format.keep_with_next=True
 direction='骨端较高' if g['gene']=='EBF1' else '肺端较高'
 p('稳定标识 '+by[g['gene']]['feature_ids']+'；本项目两来源均'+direction+'；主要证据 '+g['evidence']).paragraph_format.keep_with_next=True
 for key,label in [('biology','生物学定位'),('literature','原始研究'),('interpretation','本项目解释'),('validation','建议验证')]:p(g[key],label)

heading('如何把十四个候选组织成可检验的生物学问题')
p('第一条线索是细胞对基质环境的适应。NEBL和RTKN偏向胞质骨架／黏附，TOR1AIP1偏向核膜适应；它们可以分别检验细胞在不同机械环境中的行为。当前证据不足以把三者连成一条已确立的信号通路，也不应由迁移相关文献推导出胶原沉积增加。')
p('第二条线索是代谢和氧化还原状态。NDUFS7与复合体I有关，ZNF526的SHMT1机制指向一碳代谢和谷胱甘肽，PNKD在肾癌模型中关联乳酸与PI3K–AKT–mTOR。可把这些作为并行假设，分别测量代谢通量和功能输出；不能把它们的转录本共同升高等同于同一种代谢模式。')
p('第三条线索是间质状态及微环境相互作用。EBF1既有CAF功能研究，也有正常骨髓及肺间质背景证据。本项目中EBF1为骨端较高，提示应关注骨相关间质身份或状态，而非统一解释成肺端激活。骨肺差异还可能包含组织常驻细胞组成的变化。区分“相同状态内表达改变”与“某类细胞比例改变”，需要在来源、患者和谱系分层后分析。[23–25]')
p('GRID2、BEST1、CALHM5与三个非编码RNA适合从表达定位和分子身份切入。对这些候选，清晰、可重复的细胞归属比先拼接一条复杂通路更有价值。验证中应独立识别CAF候选，避免用目标基因本身定义细胞群后，再把同一基因的富集视为独立证据。')
heading('蛋白定位与流式研究的关系',2)
table(['候选类别','涉及基因','可采用的解释与检测定位'],[
 ('胞内蛋白','NEBL、RTKN、NDUFS7、PNKD\nTOR1AIP1、TAF10\nZNF526、EBF1','适合蛋白、胞内或原位验证；不能直接作为活细胞表面分选标志。'),
 ('潜在膜相关候选','GRID2、BEST1、CALHM5','先核对目标细胞的表面表达、表位可及性及非透膜检测；膜注释不等于已有可用CAF流式抗体。'),
 ('非编码RNA','PEF1-AS1、RGS2-AS1\nAL096854.1','适合转录本／链特异性检测与RNA原位；没有对应蛋白可直接用于表面抗体分选。')],[2.8,5.4,8.6])
heading('验证顺序和推断边界',2)
p('先验证分子身份和细胞归属，再验证患者层面的表达方向，最后开展功能干预和回补。全部14个候选均保留；按文献成熟度安排实验顺序不改变冻结名单，也不重新计算原候选排名。')
p('新增NCI数据可以检验肺端表达是否存在、来自何种细胞状态，以及患者之间是否一致。只有肺端的数据不能单独构成骨肺组间差异复现；与外部骨端比较时，应明确非配对设计、来源与部位混杂，并与已有同来源比较核对方向。技术重复和重复取材不能增加独立患者N。')
p('骨肉瘤肿瘤细胞、其他癌种CAF、正常组织间质和本项目拓展候选是不同证据层。RNA变化不等于蛋白变化，表达相关不等于因果作用，通路名称不等于功能读数。主要相反或阴性结果已保留：NEBL及TOR1AIP1存在模型间方向差异，RGS2-AS1旧名对应研究出现阴性qPCR，AL096854.1未进入其文献最终模型。[20,21,26,27]')

heading('GRID2和BEST1候选表面分选依据与预实验建议').paragraph_format.page_break_before=True
p('GRID2和BEST1具备开展活细胞表面标记预实验的依据，可优先评估其富集目标CAF或CAF亚群的能力。目前证据支持“候选分选标志物”的定位，尚未证明它们可直接用于骨肉瘤CAF分选。')
heading('GRID2的依据',2)
p('GRID2为具有胞外结构域的膜受体。针对胞外区域的AGC-039抗体已有活的、完整大鼠PC12细胞表面染色示例，支持进一步测试非透膜表面识别。该示例属于活细胞成像，并非人骨肉瘤CAF的流式验证；人源反应性、CAF表面表达量及与其他细胞的区分能力仍需实测。')
hyperlink(p('依据：'),'AGC-039胞外抗体及活细胞染色示例','https://www.thermofisher.com/antibody/product/GRID2-extracellular-Antibody-Polyclonal/AGC-039-200UL')
heading('BEST1的依据',2)
p('BEST1是细胞膜离子通道，已有针对胞外区域的ABC-001抗体，可作为表面标记预实验的试剂线索。膜定位和胞外抗体的存在尚不足以保证人CAF活细胞流式性能。应核对具体抗体的表位与人源反应性；例如常见E6-6抗体的抗原来自C端区域，不能因其识别BEST1就默认适用于活细胞表面染色。')
hyperlink(p('依据：'),'ABC-001胞外抗体资料','https://www.alomone.com/p/anti-bestrophin-1-extracellular/ABC-001')
hyperlink(p('抗体选择参考：'),'E6-6抗原区域说明','https://www.sigmaaldrich.com/US/en/product/sigma/zms1043')
heading('推荐的最小预实验',2)
p('先分别测试两支抗体。用人源目标蛋白阳性细胞和匹配阴性细胞优化抗体浓度，在活细胞、不固定、不透膜条件下检测；平行透膜组辅助判断胞内信号。设置未染、单染及必要的二抗单独对照，多色实验增加FMO设门对照，并用目标基因敲低／敲除确认信号特异性。若只有透膜后阳性，则不支持直接用于活细胞表面分选。')
p('技术条件通过后，建议先用3–5例独立患者的新鲜组织做可行性验证，优先纳入肺转移组织，逐例检测、不混池。该例数用于探索性预实验，不用于确认骨肺差异。检查组织消化是否损伤表位；在活单细胞中结合CD45、CD31排除免疫及内皮背景，并观察成纤维相关群中的GRID2或BEST1信号。')
p('在同一基础门内分选候选阳性与阴性群，回测纯度、活率和回收量，并独立核验成纤维身份及肿瘤细胞混入。骨肉瘤背景下，CD45⁻CD31⁻或FAP／PDGFR阳性不能单独证明是CAF。预实验成功的核心是：不透膜信号特异且可重复，分选后可富集经独立证据确认的目标CAF或其亚群；不要求候选覆盖全部CAF。')

heading('附表 原候选统计结果').paragraph_format.page_break_before=True
p('效应Δ为肺减骨，来自患者与部位层面汇总原始计数后计算的log2(CP10k＋0.1)表达均值差。它不是细胞均值log1p分数，也不能直接当作未经伪计数影响的表达倍数。患者标签精确置换检验按来源独立进行；两部位均有候选细胞的患者与非配对对比分开处理。')
headers=['基因','Δ S17','P S17','Δ S23','P S23']
rows=[]
for g in c['genes']:
 r=by[g['gene']];rows.append([g['gene'],f"{r['SCPCP000017__mean_lung_minus_bone']:.4f}",f"{r['SCPCP000017__P']:.6f}",f"{r['SCPCP000023__mean_lung_minus_bone']:.4f}",f"{r['SCPCP000023__P']:.6f}"])
table(headers,rows,[3.2,3.4,3.4,3.4,3.4])
q17={r['SCPCP000017__FDR'] for r in project};q23={r['SCPCP000023__FDR'] for r in project};assert len(q17)==len(q23)==1
p('S17＝SCPCP000017，骨N＝12、肺N＝3；S23＝SCPCP000023，骨N＝4、肺N＝3。效应展示至4位小数，P值展示至6位小数；原始完整精度保存在项目研究登记中。')
p('14个基因的S17 BH FDR均为 '+str(next(iter(q17)))+'；S23 BH FDR均为 '+str(next(iter(q23)))+'。BH检验家族分别为 all_cells__expanded__SCPCP000017__All_CAF__independent|genes 和 all_cells__expanded__SCPCP000023__All_CAF__independent|genes，未把不同来源、功能程序或文献结果混入同一个校正池。')
p('来源：phase84_all14_research_adoption_20261012 的 ALL14_NEXT_STAGE_RESEARCH_REGISTRY.tsv；效应定义回查phase81_allCAF_bone_lung_20261012代码。该附表是既有结果的回读，不是本轮重新检验。')

heading('参考文献与核验入口').paragraph_format.page_break_before=True
p('以下以原始研究和规范注释为主。文献年份采用出版记录；检索日期统一为2026年10月16日。补表记录仅用于说明关联线索，不替代基因功能研究。')
for r in c['references']:
 para=d.add_paragraph();para.paragraph_format.space_after=Pt(7);para.paragraph_format.line_spacing=1.12;para.paragraph_format.keep_together=True
 run=para.add_run(f"[{r['id']}] {r['title']}. {r['year']}. {r['identifier']}. ");run.font.size=Pt(9)
 hyperlink(para,'原文链接',r['url'])

heading('修订与复核说明').paragraph_format.page_break_before=True
p('本版按BIO-01至BIO-03修订，保留原版，14个候选及原统计结果不变。')
p('BIO-01｜已解决。第8节补充2018年CRISPR筛选数据分析和siRNA生长／活力干预证据，并将副标题改为“新增SHMT1研究进一步明确其代谢与抗氧化机制”。2018年关联及生长功能证据与2025年SHMT1机制证据已分开说明。[13,14]')
p('BIO-02｜已核实。第11节及参考文献[18]定位至原始Supplementary Table 3第1页第34条数据记录，PEF1-AS1符号、HR、置信区间和P值均已核对。原表无稳定ID及FDR，不能作转录本层面的认定或校正后显著性判断。正文确认最终九基因模型不含该基因。证据仅定位为乳腺癌整组织预后关联。')
p('BIO-03｜已完善。第14节明确结直肠癌CAF模型中线粒体呼吸受抑、糖酵解增强，方向与原研究一致。该结果不等同于本项目NDUFS7表达变化所代表的实际代谢通量。[23]')
p('复核范围限于上述三项及修订所影响的内容、统计表完整性和版式。文献支持的研究方向总体成立，但不能据此确认14个基因在骨肉瘤CAF中具有相同功能。')

for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
docpath=O/'骨肉瘤CAF十四个候选基因_文献调查与生物学解释_20261016_含火山图及分选建议版.docx';d.save(docpath)
read=Document(docpath);text='\n'.join(p.text for p in read.paragraphs)+'\n'+'\n'.join(cell.text for t in read.tables for row in t.rows for cell in row.cells)
assert all(g['gene'] in text for g in c['genes']);assert len(c['references'])==27
assert '500项有效检验' not in text and 'log2(CP10k＋0.1)' in text
for row in project:
 assert (row['SCPCP000017__mean_lung_minus_bone']<0)==(row['gene']=='EBF1')
 assert (row['SCPCP000023__mean_lung_minus_bone']<0)==(row['gene']=='EBF1')
(O/'BUILD_RECEIPT.json').write_text(json.dumps({'docx':str(docpath),'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'genes':14,'references':27,'tables':len(read.tables),'paragraphs':len(read.paragraphs),'render_verified':False},ensure_ascii=False,indent=2),encoding='utf-8')
print(docpath)
